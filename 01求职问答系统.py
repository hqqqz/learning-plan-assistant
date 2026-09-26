"""学习规划助手 · 单次问答 Web 应用（Streamlit）

本文件做了"入口保护"：无论你是直接运行本文件（python 01求职问答系统.py），
还是用 PyCharm 的运行配置 / streamlit run，都会以正确方式启动网页。
"""

import os
import datetime
import re

import streamlit as st
from openai import OpenAI
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx_to_pdf import convert_docx_to_pdf


def add_plan_paragraph(doc, line):
    """按层级给学习规划建议段落排版：识别一/1/（1）/- 前缀，做加粗和缩进"""
    line = line.strip()
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)

    # 一级：一、二、三…
    if re.match(r"^(一|二|三|四|五|六|七|八|九|十)、", line):
        run = p.add_run(line)
        run.bold = True
        run.font.size = Pt(12)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
    # 二级：1. 2. 3.…
    elif re.match(r"^\d+[\.、]", line):
        run = p.add_run(line)
        run.bold = True
        p.paragraph_format.left_indent = Pt(18)
    # 三级：（1）（2）…
    elif re.match(r"^（\d+）", line):
        run = p.add_run(line)
        p.paragraph_format.left_indent = Pt(36)
    # 列表项：- xxx
    elif line.startswith("- "):
        run = p.add_run(line[2:])
        p.paragraph_format.left_indent = Pt(54)
    # 普通正文（无前缀）
    else:
        run = p.add_run(line)
        p.paragraph_format.left_indent = Pt(54)
    return p


def save_report_docx(education, goal, level, time_each_day, time_target, extra, answer):
    """把学习信息和 AI 建议打包成 Word 文档，保存到本文件所在目录，返回文件路径"""
    doc = Document()

    # 设置默认中文字体，避免字体混乱
    normal = doc.styles["Normal"]
    normal.font.name = "微软雅黑"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")

    doc.add_heading("学习规划", level=0)

    doc.add_heading("一、基本信息", level=1)
    for label, val in [
        ("学历", education),
        ("目标", goal),
        ("当前基础水平", level if level.strip() else "（未填写）"),
        ("每天可投入时间", time_each_day if time_each_day.strip() else "（未填写）"),
        ("目标时间", time_target if time_target.strip() else "（未填写）"),
        ("补充信息", extra if extra.strip() else "（未填写）"),
    ]:
        doc.add_paragraph(f"{label}：{val}")

    doc.add_heading("二、学习规划建议", level=1)
    for line in answer.splitlines():
        if line.strip():  # 跳过空行，避免文档出现大量空白段
            add_plan_paragraph(doc, line)

    # 保存到本脚本所在目录，文件名带时间戳避免覆盖
    out_dir = os.path.dirname(os.path.abspath(__file__))
    filename = f"学习规划_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path


def main():
    st.set_page_config(page_title="学习规划助手", page_icon="📚", layout="centered")

    st.title("📚 学习规划助手")
    st.caption("填写基本信息，大模型为你生成对标企业要求的个性化学习规划")

    # 已生成 PDF 后，在页面顶部（回答开头）显示醒目的下载按钮
    pdf_path = st.session_state.get("pdf_path")
    if pdf_path and os.path.exists(pdf_path):
        with open(pdf_path, "rb") as f:
            pdf_bytes = f.read()
        st.download_button(
            "⬇️ 下载 PDF 学习规划",
            data=pdf_bytes,
            file_name=os.path.basename(pdf_path),
            mime="application/pdf",
            type="primary",
            use_container_width=True,
        )
        st.divider()

    # ---------------- 侧边栏：模型与密钥配置 ----------------
    with st.sidebar:
        st.header("⚙️ 模型配置")
        api_key = st.text_input(
            "DeepSeek API Key",
            type="password",
            help="在 DeepSeek 开放平台创建，形如 sk-xxxx",
        )
        base_url = st.text_input(
            "Base URL",
            value="https://api.deepseek.com",
        )
        model = st.text_input(
            "模型",
            value="deepseek-flash",
            help="deepseek-flash（V4.1，推荐）或 deepseek-v4-pro（更强）",
        )
        st.caption("没有 Key？到 DeepSeek 开放平台 platform.deepseek.com 申请。")

    # ---------------- 主区域：信息表单 ----------------
    st.write("### 请填写你的学习信息")

    with st.form("study_form", clear_on_submit=False):
        education = st.selectbox(
            "① 学历",
            ["高中及以下", "大专", "本科", "硕士", "博士", "其他"],
            index=2,
        )
        goal = st.text_input(
            "② 目标",
            placeholder="例如：转行做 AI 应用开发 / 学 Python 找到前端工作 / 提升数据分析能力…",
        )
        level = st.text_input("③ 当前的基础水平", placeholder="例如：零基础 / 学过一点 Python / 有 1 年 xx 经验…")
        time_each_day = st.text_input("④ 每天能投入的时间", placeholder="例如：每天 2 小时 / 周末集中学")
        time_target = st.text_input("⑤ 目标时间", placeholder="例如：3 个月内 / 半年内达到 xx 水平…")
        extra = st.text_area("⑥ 还有想补充的信息吗（选填）", height=80, placeholder="其他想说明的情况…")

        submitted = st.form_submit_button("🚀 获取学习规划", type="primary", use_container_width=True)

    # ---------------- 提交逻辑 ----------------
    if submitted:
        if not api_key:
            st.error("请先在左侧填入 DeepSeek API Key")
        elif not goal:
            st.warning("请填写 ②目标（必填项）")
        else:
            # 组装结构化提问
            current_time = datetime.datetime.now().strftime("%Y年%m月%d日")
            user_prompt = f"""你是一位专业、务实的私人学习规划顾问。请根据以下学习者的信息，制定一份具体、可执行的个性化学习规划。

【当前时间】现在是{current_time}。请务必以这个时间为起点来安排阶段划分、里程碑和各项学习任务的时间，不要假设或臆测其他时间点。

学习者信息：
1. 学历：{education}
2. 目标：{goal}
3. 当前基础水平：{level if level.strip() else "（未填写）"}
4. 每天可投入时间：{time_each_day if time_each_day.strip() else "（未填写）"}
5. 目标时间：{time_target if time_target.strip() else "（未填写）"}
6. 补充信息：{extra if extra.strip() else "（未填写）"}

【核心要求】请务必根据学习者的"目标"来判断其对应的岗位或方向，并结合该岗位在企业招聘时常见的任职要求（即岗位 JD 中的技能栈、工具、项目经验、能力项等）来定义学习内容，使学习规划对标企业真实的用人标准，而不是泛泛的学习建议。

请从以下几个角度给出规划（用清晰的分点结构）：
- 目标岗位的任职要求拆解（该岗位企业招聘看重什么）
- 阶段划分与时间安排（结合目标时间、每天可投入时间）
- 每个阶段的学习内容与方法（对标上述任职要求）
- 推荐的资料、工具或练习
- 检验进度的方法（里程碑）
- 常见问题与调整建议

【重要：输出格式要求】请务必用纯文本回答，方便直接生成 Word 文档：
- 不要使用任何 Markdown 标记符号（不要用 #、*、-、|、> 等符号，不要用 Markdown 表格和分隔线）
- 用"一、二、三…"作为大段落标题，用"1. 2. 3.…"和"（1）（2）"分点
- 每个要点之间用空行分隔，内容直接以普通文字呈现，不要加任何装饰符号
"""

            client = OpenAI(api_key=api_key, base_url=base_url)

            with st.chat_message("assistant"):
                st.caption(f"正在根据目标规划：{goal}")

            with st.chat_message("assistant"):
                try:
                    stream = client.chat.completions.create(
                        model=model,
                        messages=[{"role": "user", "content": user_prompt}],
                        stream=True,
                    )
                    # 流式输出，实现打字机效果
                    answer = st.write_stream(stream)
                    st.session_state["last_answer"] = answer

                    # 回答完成后，自动生成 Word 报告保存到本地
                    report_path = save_report_docx(
                        education, goal, level, time_each_day, time_target, extra, answer
                    )
                    st.success(f"✅ 已生成 Word 学习规划：\n{report_path}")

                    # 再把 Word 转成 PDF（调用本机 Word）
                    try:
                        pdf_path = convert_docx_to_pdf(report_path)
                        st.success(f"✅ 已生成 PDF 学习规划：\n{pdf_path}")
                        # 存起来并重跑页面，让顶部出现醒目的下载按钮
                        st.session_state["pdf_path"] = pdf_path
                        st.rerun()
                    except Exception as e:
                        st.warning(f"Word 转 PDF 失败：{e}（Word 文档已生成，不影响使用）")
                except Exception as e:
                    st.error(f"请求失败：{e}")
                    st.caption("请检查 API Key、Base URL 和模型名是否正确，以及网络是否可访问模型服务。")


# ---------------- 入口保护 ----------------
if __name__ == "__main__":
    from streamlit import runtime

    if runtime.exists():
        # 通过 streamlit run 启动：正常渲染页面
        main()
    else:
        # 直接 python 运行本文件：自动改用正确方式启动
        import subprocess
        import sys

        print("检测到直接运行，正在用正确方式启动网页...")
        # 跳过 Streamlit 首次运行的"填邮箱"提示，避免卡住
        env = dict(os.environ)
        env["STREAMLIT_BROWSER_GATHER_USAGE_STATS"] = "false"
        subprocess.run(
            [sys.executable, "-m", "streamlit", "run", __file__],
            input=b"\n",
            env=env,
        )
